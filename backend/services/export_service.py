"""
Export Service

Builds a multi-sheet Excel workbook from an analysis result so the computed
charts, model metrics, segments and AI insights can be imported into
Power BI (Get Data -> Excel workbook). Each sheet is a flat table with
headers in row 1, which Power BI ingests directly.
"""

from __future__ import annotations

import io
from typing import Any, Dict, List

import pandas as pd

from models.schemas import DataAnalysisData
from utils.logging_utils import get_logger

logger = get_logger(__name__)


def build_powerbi_workbook(data: DataAnalysisData) -> bytes:
    """Return the analysis result as .xlsx bytes (one sheet per table)."""
    summary = data.summary
    ml = data.ml_results or {}
    forecasting = ml.get("forecasting") or {}
    segmentation = ml.get("segmentation") or {}

    sheets: Dict[str, pd.DataFrame] = {}

    # ── AI insights & recommended questions ────────────────────────────────
    if data.insights:
        sheets["AI Insights"] = pd.DataFrame(
            {"#": range(1, len(data.insights) + 1), "Insight": data.insights}
        )
    if data.recommended_questions:
        sheets["Recommended Questions"] = pd.DataFrame(
            {"#": range(1, len(data.recommended_questions) + 1),
             "Question": data.recommended_questions}
        )

    # ── Forecasting model metrics ───────────────────────────────────────────
    models = forecasting.get("models") or {}
    if models:
        rows = [
            {"Model": name,
             "RMSE": m.get("rmse"),
             "MAE": m.get("mae"),
             "R2": m.get("r2"),
             "Best (by RMSE)": "Yes" if name == forecasting.get("best_model") else "No"}
            for name, m in models.items()
        ]
        sheets["Forecast Models"] = pd.DataFrame(rows)
        sheets["Forecast Setup"] = pd.DataFrame([
            {"Property": "Target column",  "Value": forecasting.get("target_column")},
            {"Property": "Time column",    "Value": forecasting.get("time_column")},
            {"Property": "Feature count",  "Value": forecasting.get("n_features")},
            {"Property": "Training rows",  "Value": forecasting.get("train_rows")},
            {"Property": "Test rows",      "Value": forecasting.get("test_rows")},
            {"Property": "Engine",         "Value": ml.get("engine")},
        ])
    importance = forecasting.get("feature_importance") or {}
    if importance:
        sheets["Feature Importance"] = pd.DataFrame(
            {"Feature": list(importance.keys()), "Importance": list(importance.values())}
        )

    # ── Customer segments ────────────────────────────────────────────────────
    segments = segmentation.get("segments") or []
    if segments:
        rows = []
        for seg in segments:
            row: Dict[str, Any] = {"Segment": seg.get("segment"), "Size": seg.get("size")}
            for k, v in (seg.get("profile") or {}).items():
                row[f"Avg {k}"] = v
            rows.append(row)
        sheets["Customer Segments"] = pd.DataFrame(rows)
        sheets["Segmentation Setup"] = pd.DataFrame([
            {"Property": "Method",                 "Value": segmentation.get("method")},
            {"Property": "Clusters (k)",           "Value": segmentation.get("k")},
            {"Property": "Silhouette Coefficient", "Value": segmentation.get("silhouette")},
            {"Property": "Entities clustered",     "Value": segmentation.get("n_entities")},
        ])

    # ── Chart data (same aggregations the dashboard renders) ────────────────
    ts = summary.timeseries or {}
    if ts.get("labels"):
        sheets["Timeseries"] = pd.DataFrame(
            {"Period": ts["labels"], ts.get("title", "Value"): ts.get("values", [])}
        )
    fc = summary.forecast or {}
    if fc.get("labels"):
        sheets["Forward Forecast"] = pd.DataFrame(
            {"Period": fc["labels"], fc.get("title", "Forecast"): fc.get("values", [])}
        )
    cat = summary.category_totals or {}
    if cat.get("labels"):
        sheets["Category Totals"] = pd.DataFrame(
            {"Category": cat["labels"], cat.get("title", "Value"): cat.get("values", [])}
        )
    dist = summary.distribution or {}
    if dist:
        label = summary.distribution_label or "Distribution"
        sheets["Distribution"] = pd.DataFrame(
            {label: list(dist.keys()), "Count": list(dist.values())}
        )

    # ── Summary statistics ───────────────────────────────────────────────────
    if summary.numeric_stats:
        stats_df = pd.DataFrame(summary.numeric_stats)
        stats_df.insert(0, "Statistic", stats_df.index)
        sheets["Numeric Stats"] = stats_df.reset_index(drop=True)
    if summary.top_correlations:
        sheets["Top Correlations"] = pd.DataFrame(summary.top_correlations)

    sheets["Dataset Info"] = pd.DataFrame([
        {"Property": "Filename", "Value": data.filename},
        {"Property": "Rows",     "Value": summary.rows},
        {"Property": "Columns",  "Value": summary.columns},
        {"Property": "Column names", "Value": ", ".join(summary.column_names)},
    ])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
        _embed_charts(writer.book, sheets)
    logger.info("Power BI export built: %d sheets.", len(sheets))
    return buf.getvalue()


def _embed_charts(book, sheets: Dict[str, pd.DataFrame]) -> None:
    """Add native Excel charts next to the chart-data sheets so the workbook
    shows graphs when opened (Power BI still reads the flat tables)."""
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    def _add(sheet_name: str, chart, anchor: str = "E2"):
        if sheet_name not in sheets or sheet_name not in book.sheetnames:
            return
        ws = book[sheet_name]
        n = len(sheets[sheet_name]) + 1  # +1 header row
        if n < 3:
            return
        data = Reference(ws, min_col=2, min_row=1, max_row=n)
        labels = Reference(ws, min_col=1, min_row=2, max_row=n)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        chart.height, chart.width = 9, 18
        ws.add_chart(chart, anchor)

    try:
        line = LineChart()
        line.title = sheets.get("Timeseries", pd.DataFrame()).columns[-1] if "Timeseries" in sheets else None
        _add("Timeseries", line, "D2")

        bar = BarChart()
        bar.type, bar.title = "col", (
            sheets.get("Category Totals", pd.DataFrame()).columns[-1] if "Category Totals" in sheets else None)
        _add("Category Totals", bar, "D2")

        pie = PieChart()
        pie.title = sheets.get("Distribution", pd.DataFrame()).columns[0] if "Distribution" in sheets else None
        _add("Distribution", pie, "D2")

        if "Forecast Models" in sheets:
            ws = book["Forecast Models"]
            n = len(sheets["Forecast Models"]) + 1
            mbar = BarChart()
            mbar.type, mbar.title = "col", "Model comparison (RMSE / MAE)"
            mbar.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=n),
                          titles_from_data=True)
            mbar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
            mbar.height, mbar.width = 9, 18
            ws.add_chart(mbar, "G2")
    except Exception as exc:
        logger.warning("Chart embedding failed (non-fatal, tables still exported): %s", exc)
